# Сборка DOCX с мнемокарточками для экспорта из словаря

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches
from app.schemas.mnemo import MnemoResponse

# Язык обучения
def _learning_lang_for_mnemo_row(word_lang: str | None) -> str:
    s = (word_lang or "en").strip().lower()[:2]
    return s if s in ("en", "de") else "en"

# Вставка изображения
def _image_bytes_for_docx(
    *,
    preset_assets_dir: Path,
    mnemo_cache_repo,
    user_id: int,
    word_lc: str,
    learning_lang: str,
    cached_mnemo: MnemoResponse | None,
) -> bytes | None:
    # Из БД по user + слово + язык обучения
    if mnemo_cache_repo:
        row = mnemo_cache_repo.get(user_id, word_lc, learning_lang=learning_lang)
        if row is not None and row.image_data:
            return row.image_data

    url = (cached_mnemo.image_url or "").strip() if cached_mnemo else ""
    marker = "/mnemo/preset-assets/"
    # Встроенный пресет
    if marker in url:
        fname = url.split(marker, 1)[-1].split("?", 1)[0].strip("/")
        if fname and ".." not in fname and "/" not in fname:
            p = preset_assets_dir / fname
            if p.is_file():
                return p.read_bytes()

    pub = "/mnemo/image/public/"
    # Публичная ссылка по токену
    if mnemo_cache_repo and pub in url:
        token = url.split(pub, 1)[-1].split("?", 1)[0].strip("/")
        if token:
            row2 = mnemo_cache_repo.get_by_image_token(token)
            if row2 is not None and row2.image_data:
                return row2.image_data

    return None

# Собираем DOCX по списку id записей словаря пользователя
def build_user_cards_docx(
    *,
    user_id: int,
    entry_ids: list[int],
    user_words_repo,
    mnemo_cache_repo,
    decomposition_cache_repo,
    mnemo_service,
    preset_assets_dir: Path,
) -> bytes:
    doc = Document()
    doc.add_heading("MnemoCards", level=0)

    seen: set[int] = set()
    added = 0

    for eid in entry_ids:
        # Пропуск дубликатов и невалидных id
        if eid <= 0 or eid in seen:
            continue
        seen.add(eid)
        entry = user_words_repo.get(eid, user_id)
        if entry is None:
            continue

        added += 1
        w = (entry.word or "").strip()
        word_lc = w.lower()
        trans = (entry.translation or "").strip()
        transcr = (entry.transcription or "").strip()
        cat = (entry.category or "").strip()

        ll = _learning_lang_for_mnemo_row(getattr(entry, "word_language", None))

        # Мнемоника из сервиса (память/файловый кэш) и/или из репозитория БД
        cached_mnemo: MnemoResponse | None = None
        if mnemo_service is not None:
            cached_mnemo = mnemo_service.get_cache(w, user_id=user_id, learning_lang=ll)

        row = mnemo_cache_repo.get(user_id, word_lc, learning_lang=ll) if mnemo_cache_repo else None
        phrase = ""
        if cached_mnemo and (cached_mnemo.mnemonic_phrase_ru or "").strip():
            phrase = (cached_mnemo.mnemonic_phrase_ru or "").strip()
        elif row and (row.mnemonic_phrase_ru or "").strip():
            phrase = (row.mnemonic_phrase_ru or "").strip()

        parts = decomposition_cache_repo.get(entry.id) if decomposition_cache_repo else None

        # Одна карточка = заголовок-слово + метаданные + стих + картинка
        doc.add_heading(w, level=1)
        doc.add_paragraph(f"Перевод: {trans}")
        if transcr:
            doc.add_paragraph(f"Транскрипция (IPA): {transcr}")
        if cat:
            doc.add_paragraph(f"Категория: {cat}")
        if parts:
            doc.add_paragraph("Декомпозиция: " + " + ".join(parts))

        if phrase:
            doc.add_paragraph("Мнемостих:")
            for line in phrase.splitlines():
                ln = line.strip()
                if ln:
                    doc.add_paragraph(ln)
        else:
            doc.add_paragraph("(Мнемостих не найден в сохранённом кэше для этого слова)")

        img = _image_bytes_for_docx(
            preset_assets_dir=preset_assets_dir,
            mnemo_cache_repo=mnemo_cache_repo,
            user_id=user_id,
            word_lc=word_lc,
            learning_lang=ll,
            cached_mnemo=cached_mnemo,
        )
        if img and len(img) > 10:
            bio = BytesIO(img)
            bio.seek(0)
            try:
                doc.add_picture(bio, width=Inches(4.0))
            except Exception:
                doc.add_paragraph("(Изображение не удалось вставить в DOCX)")

        doc.add_paragraph("")

    if added == 0:
        doc.add_paragraph("Нет записей для экспорта: проверьте id и авторизацию")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
